import io
from xml.sax.saxutils import escape

from utils.config import EXPORT_DIR

_MAP = {"\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"', "\u2013": "-", "\u2014": "-", "\u2026": "...", "\u2192": "->", "\u00a0": " "}


def _lat(s):
    s = str(s or "")
    for k, v in _MAP.items():
        s = s.replace(k, v)
    return s.encode("latin-1", "replace").decode("latin-1")


def _ref_line(r):
    parts = [r.get("authors"), f"({r['year']})" if r.get("year") else "", r.get("title"), r.get("venue")]
    return ". ".join(p for p in parts if p)


def _how_lines(h):
    s, e, v = h["search"], h["evaluation"], h["verification"]
    lines = [
        f"{s['completed']} of {s['searches']} searches completed ({s['failed']} failed); {s['raw']} raw results, {s.get('unique', 0)} unique sources.",
        f"{e['kept']} sources evaluated and kept; {e['no_text']} dropped for lacking readable text, {e['low_score']} below the credibility threshold.",
        f"{v['candidates']} candidate claims checked: {v['verified']} verified by multiple sources, {v['partial']} partially verified; "
        f"{v['removed_unverified']} removed because the quote was not found in the source, {v['removed_lowcred']} removed as single low-credibility sources.",
        (f"Conflict check run: {h['conflicts_found']} genuine disagreement(s)." if h["conflicts_checked"] else "Conflict check not run."),
    ]
    return lines + ["Limitation: " + x for x in h["limitations"]]


def to_pdf(task, opts, save=True):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    r = task["report"]
    refs = {x["n"]: x for x in r["references"]}
    ss = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=ss["BodyText"], fontSize=10, leading=14)
    small = ParagraphStyle("s", parent=body, fontSize=8.5, leading=11, textColor=colors.HexColor("#555555"))
    h1, h2 = ss["Title"], ss["Heading2"]

    def P(t, st=body):
        return Paragraph(t, st)

    def cites(ns):
        out = []
        for n in ns:
            u = refs.get(n, {}).get("url", "")
            out.append(f'<link href="{escape(u, {chr(34): "&quot;"})}" color="#2457d6">[{n}]</link>' if u else f"[{n}]")
        return " ".join(out)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm, title=_lat(task["topic"]))
    st = [P(escape(_lat(task["topic"])), h1), P(escape(f"{task['depth']} · {task['created']} · {_lat(task['model'])}"), small), Spacer(1, 6)]
    st += [P("Introduction", h2), P(escape(_lat(r["introduction"])))]
    st.append(P("Key Findings", h2))
    for f in r["findings"]:
        st.append(P(f"<b>{f['n']}. Claim:</b> {escape(_lat(f['claim']))} {cites([f['source_n']] + f['support_ns'])}"))
        if opts.get("export_evidence", True):
            st.append(P(f"<b>Evidence:</b> \"{escape(_lat(f['evidence']))}\""))
        src = refs.get(f["source_n"], {})
        st.append(P(f"<b>Source:</b> {escape(_lat(src.get('title', '')))} {cites([f['source_n']])}"))
        st.append(P(f"<b>Verification:</b> {escape(f['status'])} - {escape(_lat(f['detail']))} Confidence: {f['confidence']}."))
        if f["evidence_type"] == "correlational":
            st.append(P("Correlational evidence - does not establish causation.", small))
        if f["caveat"]:
            st.append(P(f"Caveat: {escape(_lat(f['caveat']))}", small))
        st.append(Spacer(1, 6))
    if r["perspectives"]:
        st.append(P("Different Perspectives", h2))
        for p in r["perspectives"]:
            st.append(P(f"<b>{escape(_lat(p['topic']))}</b> - {escape(_lat(p['explanation']))}"))
            for x in p["positions"]:
                st.append(P(f"&bull; F{x['n']}: {escape(_lat(x['claim']))} {cites(x['cites'])}"))
    if r["statistics"]:
        st.append(P("Important Statistics", h2))
        for x in r["statistics"]:
            st.append(P(f"<b>{escape(_lat(x['value']))}</b> - {escape(_lat(x['context']))} {cites(x['cites'])} ({x['status']})"))
    if opts.get("export_map", True) and r["findings"]:
        st.append(P("Evidence Map", h2))
        rows = [[P("<b>#</b>", small), P("<b>Claim</b>", small), P("<b>Sources</b>", small), P("<b>Verification</b>", small)]]
        for f in r["findings"]:
            rows.append([P(str(f["n"]), small), P(escape(_lat(f["claim"])), small), P(cites([f["source_n"]] + f["support_ns"]), small),
                         P(escape(f"{f['status']} / {f['confidence']}"), small)])
        t = Table(rows, colWidths=[9 * mm, 85 * mm, 35 * mm, 43 * mm], repeatRows=1)
        t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cccccc")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                               ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef1f6"))]))
        st.append(t)
    st.append(P("How We Checked", h2))
    for line in _how_lines(r["how"]):
        st.append(P("&bull; " + escape(_lat(line))))
    if opts.get("export_log"):
        st.append(P("Process Log", h2))
        for line in task.get("log", []):
            st.append(P(escape(_lat(line)), small))
    st += [P("Conclusion", h2), P(escape(_lat(r["conclusion"]))), P("References", h2)]
    for x in r["references"]:
        u = x["url"]
        st.append(P(f"[{x['n']}] {escape(_lat(_ref_line(x)))} " + (f'<link href="{escape(u, {chr(34): "&quot;"})}" color="#2457d6">{escape(u)}</link>' if u else ""), small))
    doc.build(st)
    data = buf.getvalue()
    if save:
        (EXPORT_DIR / f"{task['id']}.pdf").write_bytes(data)
    return data


def _docx_link(par, text, url):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    r_id = par.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), "2457D6")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(c)
    rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    h.append(run)
    par._p.append(h)


def to_docx(task, opts, save=True):
    from docx import Document
    from docx.shared import Pt

    r = task["report"]
    refs = {x["n"]: x for x in r["references"]}
    d = Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(10.5)

    def cites(par, ns):
        for n in ns:
            u = refs.get(n, {}).get("url", "")
            par.add_run(" ")
            _docx_link(par, f"[{n}]", u) if u else par.add_run(f"[{n}]")

    def para(label, text, ns=None, italic=False):
        p = d.add_paragraph()
        if label:
            p.add_run(label + " ").bold = True
        run = p.add_run(text)
        run.italic = italic
        if ns:
            cites(p, ns)
        return p

    d.add_heading(task["topic"], 0)
    d.add_paragraph(f"{task['depth']} · {task['created']} · {task['model']}").runs[0].italic = True
    d.add_heading("Introduction", 1)
    d.add_paragraph(r["introduction"])
    d.add_heading("Key Findings", 1)
    for f in r["findings"]:
        para(f"{f['n']}. Claim:", f["claim"], [f["source_n"]] + f["support_ns"])
        if opts.get("export_evidence", True):
            para("Evidence:", f"“{f['evidence']}”")
        src = refs.get(f["source_n"], {})
        para("Source:", src.get("title", ""), [f["source_n"]])
        para("Verification:", f"{f['status']} — {f['detail']} Confidence: {f['confidence']}.")
        if f["evidence_type"] == "correlational":
            para("", "Correlational evidence — does not establish causation.", italic=True)
        if f["caveat"]:
            para("Caveat:", f["caveat"], italic=True)
    if r["perspectives"]:
        d.add_heading("Different Perspectives", 1)
        for p in r["perspectives"]:
            para(p["topic"] + ":", p["explanation"])
            for x in p["positions"]:
                para("", f"F{x['n']}: {x['claim']}", x["cites"])
    if r["statistics"]:
        d.add_heading("Important Statistics", 1)
        for x in r["statistics"]:
            para(x["value"] + " —", f"{x['context']} ({x['status']})", x["cites"])
    if opts.get("export_map", True) and r["findings"]:
        d.add_heading("Evidence Map", 1)
        t = d.add_table(rows=1, cols=4)
        t.style = "Table Grid"
        for i, h in enumerate(["#", "Claim", "Sources", "Verification / confidence"]):
            t.rows[0].cells[i].text = h
        for f in r["findings"]:
            c = t.add_row().cells
            c[0].text, c[1].text = str(f["n"]), f["claim"]
            c[2].text = ", ".join(f"[{n}]" for n in [f["source_n"]] + f["support_ns"])
            c[3].text = f"{f['status']} / {f['confidence']}"
    d.add_heading("How We Checked", 1)
    for line in _how_lines(r["how"]):
        d.add_paragraph(line, style="List Bullet")
    if opts.get("export_log"):
        d.add_heading("Process Log", 1)
        for line in task.get("log", []):
            d.add_paragraph(line)
    d.add_heading("Conclusion", 1)
    d.add_paragraph(r["conclusion"])
    d.add_heading("References", 1)
    for x in r["references"]:
        p = d.add_paragraph(f"[{x['n']}] {_ref_line(x)} ")
        if x["url"]:
            _docx_link(p, x["url"], x["url"])
    buf = io.BytesIO()
    d.save(buf)
    data = buf.getvalue()
    if save:
        (EXPORT_DIR / f"{task['id']}.docx").write_bytes(data)
    return data
